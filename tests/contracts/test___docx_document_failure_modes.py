"""F68 (ADR-024 Bundle A) — failure-mode contract for ``_DocxDocument``.

``_DocxDocument`` is the ``docx.Document`` wire shape the ``docx``
extractor reads (``paragraphs`` / ``tables`` / ``element`` /
``core_properties``). The failure shapes that matter to the connector
pipeline calling ``DocxExtractor.extract``:

* an empty body (no paragraphs) — an empty, zero-confidence document that
  fails ``quality_ok`` so the escalation chain routes it onward;
* no tables — the markdown carries the paragraphs only, never an empty
  table skeleton;
* a corrupt package whose ``document.xml`` lost its ``<w:body>`` — the
  extractor must degrade to an empty document with no tracked changes,
  not raise ``AttributeError`` out of the walk;
* blank core properties — metadata degrades to ``None`` (never ``""``).

One body per method, two document implementations (F43 limb 2), both
driven through the real :class:`DocxExtractor`:

* real ``python-docx`` documents built in memory (and, for the corrupt
  case, re-zipped without ``<w:body>``) opened via the default opener;
* :class:`tests.fakes.FakeDocxDocument` injected through ``document_opener``.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable
from types import SimpleNamespace
from typing import Any

import docx
import pytest

from kairix.extractors.docx import DocxExtractor
from kairix.extractors.docx import version as docx_version
from tests.fakes import FakeDocxDocument

pytestmark = pytest.mark.contract

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_HEADING = "Release checklist"
_BODY = "Roll the worker first, then the MCP server."

_Case = tuple[DocxExtractor, bytes]


def _docx_bytes(build: Callable[[Any], None]) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _with_heading_and_body(document: Any) -> None:
    document.add_heading(_HEADING, level=1)
    document.add_paragraph(_BODY)


def _with_blank_core_properties(document: Any) -> None:
    document.core_properties.title = "   "
    document.core_properties.author = ""


def _without_body(raw: bytes) -> bytes:
    """Re-zip a docx with ``<w:body>`` stripped from ``word/document.xml``."""
    source = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                data = re.sub(rb"<w:body>.*</w:body>", b"", data, flags=re.S)
            target.writestr(info, data)
    return out.getvalue()


def _real(raw: bytes) -> _Case:
    return DocxExtractor(version=docx_version), raw


def _fake(document: FakeDocxDocument) -> _Case:
    extractor = DocxExtractor(version=docx_version, document_opener=lambda: lambda _path: document)
    return extractor, b"PK\x03\x04fake-docx"


def _paragraph(text: str, style: str) -> Any:
    return SimpleNamespace(text=text, style=SimpleNamespace(name=style))


_EMPTY_BODY: list[Callable[[], _Case]] = [
    lambda: _real(_docx_bytes(lambda _document: None)),
    lambda: _fake(FakeDocxDocument(paragraphs=[])),
]
_NO_TABLES: list[Callable[[], _Case]] = [
    lambda: _real(_docx_bytes(_with_heading_and_body)),
    lambda: _fake(
        FakeDocxDocument(paragraphs=[_paragraph(_HEADING, "Heading 1"), _paragraph(_BODY, "Normal")], tables=[])
    ),
]
_BODYLESS: list[Callable[[], _Case]] = [
    lambda: _real(_without_body(_docx_bytes(_with_heading_and_body))),
    lambda: _fake(FakeDocxDocument(element=SimpleNamespace(body=None))),
]
_BLANK_PROPS: list[Callable[[], _Case]] = [
    lambda: _real(_docx_bytes(_with_blank_core_properties)),
    lambda: _fake(FakeDocxDocument(core_properties=SimpleNamespace(title="   ", author="", created=None))),
]
_IDS = ["real-python-docx", "fake"]


@pytest.mark.parametrize("factory", _EMPTY_BODY, ids=_IDS)
def test_paragraphs_returns_empty_document_has_no_content_and_escalates(factory: Callable[[], _Case]) -> None:
    """No paragraphs: empty markdown, zero confidence, ``quality_ok`` False.

    Sabotage proof: in ``DocxExtractor.quality_ok`` change the
    ``if len(doc.markdown) < _QUALITY_MIN_CHARS: return False`` guard to
    ``return True``. Re-run: both cases fail on the quality assertion.
    Restored.
    """
    extractor, raw = factory()

    doc = extractor.extract(raw, _DOCX_MIME)

    assert (doc.markdown, doc.confidence) == ("", 0.0)
    assert extractor.quality_ok(doc) is False


@pytest.mark.parametrize("factory", _NO_TABLES, ids=_IDS)
def test_tables_returns_empty_markdown_carries_paragraphs_only(factory: Callable[[], _Case]) -> None:
    """No tables: the markdown is exactly the rendered paragraphs.

    Sabotage proof: in ``_render_document`` append
    ``lines.append(_table_to_markdown([[""]]))`` after the table loop.
    Re-run: both cases fail because an empty table skeleton is appended.
    Restored.
    """
    extractor, raw = factory()

    doc = extractor.extract(raw, _DOCX_MIME)

    assert doc.markdown == f"# {_HEADING}\n\n{_BODY}"
    assert "|" not in doc.markdown


@pytest.mark.parametrize("factory", _BODYLESS, ids=_IDS)
def test_element_returns_partial_bodyless_package_degrades_to_empty_document(factory: Callable[[], _Case]) -> None:
    """A document element with no ``<w:body>`` yields an empty document with
    no tracked changes — not an ``AttributeError`` from the body walk.

    Sabotage proof: in ``_safe_body_xml`` replace ``if body is None:
    return ""`` with ``return body.xml``. Re-run: both cases fail with
    ``AttributeError`` out of ``extract``. Restored.
    """
    extractor, raw = factory()

    doc = extractor.extract(raw, _DOCX_MIME)

    assert doc.markdown == ""
    assert extractor.last_extract_had_tracked_changes is False
    assert extractor.quality_ok(doc) is False


@pytest.mark.parametrize("factory", _BLANK_PROPS, ids=_IDS)
def test_core_properties_returns_empty_metadata_degrades_to_none(factory: Callable[[], _Case]) -> None:
    """Blank title / author core properties become ``None``, not ``""``.

    Sabotage proof: in ``kairix/extractors/docx/extractor.py::_clean_string``
    change ``return stripped or None`` to ``return stripped``. Re-run: both
    cases fail because title / author come back as ``""``. Restored.
    """
    extractor, raw = factory()

    meta = extractor.extract(raw, _DOCX_MIME).metadata

    assert (meta.title, meta.author) == (None, None)
